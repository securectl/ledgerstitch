# SPDX-License-Identifier: AGPL-3.0-only
"""CPA-ready Excel workbook (openpyxl).

Tabs: Summary, Schedule C, Business, Personal, Needs Review, Excluded,
All Transactions, By Account.  Totals on the detail tabs are live
``=SUBTOTAL(9, ...)`` formulas; the Summary also carries static values computed
by LedgerStitch (openpyxl cannot write cached formula results, and some viewers
never recalculate), clearly labelled as such.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from ..models import Account, Direction, Lane, Source, Transaction
from .summary import (DISCLAIMER, EXPENSE_LANES, MEALS_LINE, MEALS_NOTE, accounts_by_id,
                      category_of, expense_value, income_value, q2, schedule_c_rows,
                      split_active, summarize)

ACCENT = "1F6F5C"
ACCENT_LIGHT = "E3F0EC"
MONEY = '_($* #,##0.00_);[Red]_($* (#,##0.00);_($* "-"??_);_(@_)'
PCT = "0%"
DATE_FMT = "yyyy-mm-dd"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
TITLE_FONT = Font(bold=True, size=14, color=ACCENT)
BOLD = Font(bold=True)
NOTE_FONT = Font(italic=True, color="555555", size=9)
LINK_FONT = Font(color="0563C1", underline="single")
TOTAL_FILL = PatternFill("solid", fgColor=ACCENT_LIGHT)
TOP_BORDER = Border(top=Side(style="thin", color=ACCENT))

SOURCE_LABEL = {Source.RULE: "Rule", Source.AI: "AI", Source.USER: "User",
                Source.HEURISTIC: "Heuristic", Source.NONE: "None"}
LANE_LABEL = {Lane.BUSINESS: "Business", Lane.PERSONAL: "Personal", Lane.TRANSFER: "Transfer",
              Lane.INCOME: "Income", Lane.REVIEW: "Needs review"}

ALL_COLUMNS = ["Date", "Account", "Description", "Merchant", "Category", "Lane", "Sched C",
               "Amount", "Status", "Source", "Note", "Txn ID"]
ALL_WIDTHS = [12, 24, 44, 22, 26, 13, 9, 15, 10, 11, 28, 18]
LANE_COLUMNS = ["Date", "Account", "Description", "Merchant", "Category", "Sched C",
                "Expense", "Status", "Source", "Note", "Txn ID"]
LANE_WIDTHS = [12, 24, 44, 22, 26, 9, 15, 10, 11, 28, 18]
EXCLUDED_COLUMNS = ["Reason", "Date", "Account", "Description", "Category", "Lane",
                    "Amount", "Txn ID", "Detail"]
EXCLUDED_WIDTHS = [22, 12, 24, 44, 26, 13, 15, 18, 36]


# --------------------------------------------------------------------------- cell helpers
def _text(ws, row: int, col: int, value, **style):
    """Write a string that must never be interpreted as a formula."""
    c = ws.cell(row=row, column=col, value="" if value is None else str(value))
    c.data_type = "s"
    for k, v in style.items():
        setattr(c, k, v)
    return c


def _money(ws, row: int, col: int, value, bold: bool = False):
    v = float(q2(value)) if isinstance(value, Decimal) else value
    c = ws.cell(row=row, column=col, value=v)
    c.number_format = MONEY
    if bold:
        c.font = BOLD
    return c


def _header(ws, row: int, columns: list[str], widths: list[int] | None = None):
    for i, name in enumerate(columns, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=False)
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]


def _total_row(ws, row: int, ncols: int, label_col: int, label: str, sum_cols: list[int],
               first: int, last: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = TOTAL_FILL
        cell.border = TOP_BORDER
    _text(ws, row, label_col, label, font=BOLD)
    for col in sum_cols:
        L = get_column_letter(col)
        formula = f"=SUBTOTAL(9,{L}{first}:{L}{last})" if last >= first else "=0"
        c = ws.cell(row=row, column=col, value=formula)
        c.number_format = MONEY
        c.font = BOLD


def _link(cell, sheet: str, row: int):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A{row}",
                               display=str(cell.value))
    cell.font = LINK_FONT


def _acct_label(accts: dict[str, Account], aid: str) -> str:
    a = accts.get(aid)
    return a.label if a else aid


# --------------------------------------------------------------------------- sheets
def _write_all(ws, active: list[Transaction], accts) -> tuple[dict[str, int], int]:
    """All Transactions tab. Returns ({category: first row}, total row)."""
    _header(ws, 1, ALL_COLUMNS, ALL_WIDTHS)
    first_row: dict[str, int] = {}
    r = 1
    for r, t in enumerate(active, 2):
        cat = category_of(t)
        first_row.setdefault((t.lane.value, cat), r)
        first_row.setdefault(cat, r)
        ws.cell(row=r, column=1, value=t.date).number_format = DATE_FMT
        _text(ws, r, 2, _acct_label(accts, t.account_id))
        _text(ws, r, 3, t.description)
        _text(ws, r, 4, t.merchant)
        _text(ws, r, 5, cat)
        _text(ws, r, 6, LANE_LABEL[t.lane])
        _text(ws, r, 7, t.schedule_c or "")
        _money(ws, r, 8, t.signed_amount)
        _text(ws, r, 9, t.status.value.capitalize())
        _text(ws, r, 10, SOURCE_LABEL.get(t.source, t.source.value))
        _text(ws, r, 11, t.note)
        _text(ws, r, 12, t.id)
    last = max(r, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ALL_COLUMNS))}{max(last, 2)}"
    ws.freeze_panes = "A2"
    total = last + 2
    _total_row(ws, total, len(ALL_COLUMNS), 7, "Visible total", [8], 2, last)
    _text(ws, total + 1, 1, "Amount is signed: money out is negative, money in is positive. "
          "The total follows the AutoFilter (SUBTOTAL).", font=NOTE_FONT)
    return first_row, total


def _write_lane(ws, rows: list[Transaction], accts, title: str) -> int:
    """Business / Personal / Needs Review tab. Returns the total row number."""
    _header(ws, 1, LANE_COLUMNS, LANE_WIDTHS)
    r = 1
    for r, t in enumerate(rows, 2):
        ws.cell(row=r, column=1, value=t.date).number_format = DATE_FMT
        _text(ws, r, 2, _acct_label(accts, t.account_id))
        _text(ws, r, 3, t.description)
        _text(ws, r, 4, t.merchant)
        _text(ws, r, 5, category_of(t))
        _text(ws, r, 6, t.schedule_c or "")
        _money(ws, r, 7, expense_value(t))
        _text(ws, r, 8, t.status.value.capitalize())
        _text(ws, r, 9, SOURCE_LABEL.get(t.source, t.source.value))
        _text(ws, r, 10, t.note)
        _text(ws, r, 11, t.id)
    last = r
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LANE_COLUMNS))}{max(last, 2)}"
    ws.freeze_panes = "A2"
    total = last + 2
    _total_row(ws, total, len(LANE_COLUMNS), 6, f"{title} total", [7], 2, last)
    _text(ws, total + 1, 1, "Expense: purchases positive, refunds negative (net).", font=NOTE_FONT)
    return total


def _excluded_reason(t: Transaction) -> tuple[str, str]:
    if t.duplicate_of is not None:
        return "Duplicate", f"Duplicate of {t.duplicate_of}"
    if t.lane == Lane.INCOME:
        return "Income (not an expense)", "Reported separately as income"
    if t.transfer_group:
        return "Transfer / card payment", f"Matched pair {t.transfer_group}"
    if t.refund_of:
        return "Refund", f"Refund of {t.refund_of}"
    return "Transfer / card payment", "Own-account movement (other leg not imported)"


def _write_excluded(ws, rows: list[Transaction], accts):
    _header(ws, 1, EXCLUDED_COLUMNS, EXCLUDED_WIDTHS)
    for r, t in enumerate(rows, 2):
        reason, detail = _excluded_reason(t)
        _text(ws, r, 1, reason)
        ws.cell(row=r, column=2, value=t.date).number_format = DATE_FMT
        _text(ws, r, 3, _acct_label(accts, t.account_id))
        _text(ws, r, 4, t.description)
        _text(ws, r, 5, category_of(t))
        _text(ws, r, 6, LANE_LABEL[t.lane])
        _money(ws, r, 7, t.signed_amount)
        _text(ws, r, 8, t.id)
        _text(ws, r, 9, detail)
    last = max(len(rows) + 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCLUDED_COLUMNS))}{last}"
    ws.freeze_panes = "A2"
    _text(ws, last + 2, 1, "These rows are NOT counted in expense totals. Duplicates are excluded "
          "from every other tab and export.", font=NOTE_FONT)


def _write_schedule_c(ws, active: list[Transaction], business_name: str, year: int):
    _text(ws, 1, 1, f"Schedule C (Form 1040) - Part II expenses - {year}", font=TITLE_FONT)
    _text(ws, 2, 1, business_name or "")
    _text(ws, 3, 1, DISCLAIMER, font=NOTE_FONT)
    cols = ["Line", "Description", "Gross", "Deductible %", "Deductible estimate", "Txns"]
    _header(ws, 5, cols, [10, 46, 16, 13, 20, 8])
    rows = schedule_c_rows(active)
    r = 5
    for r, row in enumerate(rows, 6):
        _text(ws, r, 1, row.line)
        _text(ws, r, 2, row.name)
        _money(ws, r, 3, row.gross)
        c = ws.cell(row=r, column=4, value=float(row.pct))
        c.number_format = PCT
        _money(ws, r, 5, row.deductible)
        ws.cell(row=r, column=6, value=row.count)
    last = r
    total = last + 1
    _total_row(ws, total, len(cols), 2, "Total (business lane)", [3, 5], 6, last)
    n = total + 2
    _text(ws, n, 1, f"Note: {MEALS_NOTE}", font=NOTE_FONT)
    _text(ws, n + 1, 1, "Business lane only; refunds are netted. Line-level detail is on the "
          "Business tab. Gross and deductible values are computed by LedgerStitch.", font=NOTE_FONT)
    _text(ws, n + 2, 1, DISCLAIMER, font=BOLD)
    ws.freeze_panes = "A6"


def _write_by_account(ws, active: list[Transaction], accts):
    cols = ["Account", "Institution", "Type", "Business", "Personal", "Needs review",
            "Transfers out", "Transfers in", "Income", "Txns"]
    _header(ws, 1, cols, [28, 16, 13, 15, 15, 15, 15, 15, 15, 8])
    agg: dict[str, dict] = defaultdict(lambda: defaultdict(lambda: Decimal("0")))
    counts: dict[str, int] = defaultdict(int)
    for t in active:
        a = agg[t.account_id]
        counts[t.account_id] += 1
        if t.lane in EXPENSE_LANES:
            a[t.lane.value] += expense_value(t)
        elif t.lane == Lane.TRANSFER:
            a["tout" if t.direction == Direction.OUT else "tin"] += t.amount
        elif t.lane == Lane.INCOME:
            a["income"] += income_value(t)
    order = list(accts) + sorted(set(agg) - set(accts))
    r = 1
    for aid in order:
        if aid not in agg:
            continue
        r += 1
        acct = accts.get(aid)
        _text(ws, r, 1, _acct_label(accts, aid))
        _text(ws, r, 2, acct.institution if acct else "")
        _text(ws, r, 3, acct.kind.value.replace("_", " ").title() if acct else "")
        for col, key in enumerate(("business", "personal", "review", "tout", "tin", "income"), 4):
            _money(ws, r, col, agg[aid][key])
        ws.cell(row=r, column=10, value=counts[aid])
    last = r
    _total_row(ws, last + 1, len(cols), 1, "All accounts", [4, 5, 6, 7, 8, 9], 2, last)
    _text(ws, last + 3, 1, "Expense columns are net of refunds. Transfers and income are not expenses.",
          font=NOTE_FONT)
    ws.freeze_panes = "A2"


def _write_summary(ws, s: dict, business_name: str, year: int, first_row: dict,
                   lane_totals_cells: dict[str, str]):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 18
    _text(ws, 1, 1, business_name or "Expense report", font=TITLE_FONT)
    _text(ws, 2, 1, "Tax year")
    ws.cell(row=2, column=2, value=year).alignment = Alignment(horizontal="left")
    _text(ws, 3, 1, "Generated")
    c = ws.cell(row=3, column=2, value=date.today())
    c.number_format = DATE_FMT
    c.alignment = Alignment(horizontal="left")
    _text(ws, 4, 1, DISCLAIMER, font=NOTE_FONT)

    _header(ws, 6, ["Lane totals", "Computed by LedgerStitch", "Live formula check"])
    t = s["totals"]
    lanes = [("Business expenses", "business"), ("Personal spending", "personal"),
             ("Needs review (unresolved)", "review"), ("Income (not an expense)", "income"),
             ("Transfers out (excluded)", "transfers_out"), ("Transfers in (excluded)", "transfers_in")]
    r = 6
    for r, (label, key) in enumerate(lanes, 7):
        _text(ws, r, 1, label)
        _money(ws, r, 2, t[key])
        if key in lane_totals_cells:
            c = ws.cell(row=r, column=3, value=f"={lane_totals_cells[key]}")
            c.number_format = MONEY
    r += 1
    _text(ws, r, 1, "Total expenses (business + personal + review)", font=BOLD)
    _money(ws, r, 2, t["expenses"], bold=True)
    r += 1
    _text(ws, r, 1, "Static values are computed by LedgerStitch in Python; the formula column "
          "recalculates from the detail tabs when opened in Excel.", font=NOTE_FONT)
    r += 1
    counts = s["counts"]
    _text(ws, r, 1, f"{counts['total']} transactions; {counts['duplicates']} duplicates removed; "
          f"{s['review_count']} need review; {s['uncategorized_count']} uncategorized.", font=NOTE_FONT)

    r += 2
    _header(ws, r, ["Category (click to open)", "Lane", "Sched C", "Txns", "Net total"])
    for cat in s["by_category"]:
        r += 1
        c = _text(ws, r, 1, cat["category"])
        target = first_row.get((cat["lane"], cat["category"])) or first_row.get(cat["category"])
        if target:
            _link(c, "All Transactions", target)
        _text(ws, r, 2, LANE_LABEL[Lane(cat["lane"])])
        _text(ws, r, 3, cat["schedule_c"] or "")
        ws.cell(row=r, column=4, value=cat["count"])
        _money(ws, r, 5, cat["total"])
    ws.freeze_panes = "A6"


# --------------------------------------------------------------------------- entry point
SHEETS = ["Summary", "Schedule C", "Business", "Personal", "Needs Review", "Excluded",
          "All Transactions", "By Account"]


def export_excel(txns, accounts, path: Path, business_name: str, year: int) -> Path:
    """Write the CPA workbook for ``year`` (``None`` = all dates) and return ``path``."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    txns = list(txns)
    accts = accounts_by_id(accounts)
    active, dups = split_active(txns, year)
    s = summarize(txns, year)

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    sheets = {"Summary": ws_sum}
    for name in SHEETS[1:]:
        sheets[name] = wb.create_sheet(name)
    for ws in sheets.values():
        ws.sheet_properties.tabColor = ACCENT

    first_row, _ = _write_all(sheets["All Transactions"], active, accts)
    _write_schedule_c(sheets["Schedule C"], active, business_name, year)
    lane_cells: dict[str, str] = {}
    for sheet, lane, title in (("Business", Lane.BUSINESS, "Business"),
                               ("Personal", Lane.PERSONAL, "Personal"),
                               ("Needs Review", Lane.REVIEW, "Needs review")):
        rows = [t for t in active if t.lane == lane]
        total = _write_lane(sheets[sheet], rows, accts, title)
        lane_cells[lane.value] = f"'{sheet}'!G{total}"
    excluded = [t for t in active if t.lane in (Lane.TRANSFER, Lane.INCOME)] + dups
    excluded.sort(key=lambda t: (t.duplicate_of is not None, t.date, t.account_id, t.seq))
    _write_excluded(sheets["Excluded"], excluded, accts)
    _write_by_account(sheets["By Account"], active, accts)
    _write_summary(ws_sum, s, business_name, year, first_row, lane_cells)

    wb.properties.title = f"{business_name or 'LedgerStitch'} - {year} expense report"
    wb.properties.creator = "LedgerStitch"
    wb.save(path)
    return path
