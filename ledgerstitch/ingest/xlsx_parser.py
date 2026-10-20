# SPDX-License-Identifier: AGPL-3.0-only
"""Excel (.xlsx) statement exports.

An .xlsx is converted to CSV text in memory and handed to the normal CSV
pipeline (`detect` + `csv_parser`), so every CSV feature (known profiles,
generic column inference, learned formats, debit/credit split, footers, ...)
works unchanged for spreadsheets.

* The first sheet with a header-like row (≥3 text cells followed within 5
  rows by a row holding a date or a number) is used; earlier blank/summary
  sheets are skipped. Rows above the header (preamble) are kept so header
  location works exactly as it does for CSVs.
* Date cells become ISO ``YYYY-MM-DD``; floats are written without binary
  noise (``0.30000000000000004`` → ``0.3``).
* openpyxl never executes macros or formulas; cached formula values are read
  (``data_only=True``). Zip bombs are refused (entry count / declared size).
* Legacy binary ``.xls`` is rejected with a clear message.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from typing import Optional

__all__ = ["sniff_xlsx", "is_legacy_xls", "xlsx_to_csv_bytes", "parse_xlsx", "detect_xlsx"]

MAX_ENTRIES = 2000
MAX_UNCOMPRESSED = 200 * 1024 * 1024
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def is_legacy_xls(data: bytes) -> bool:
    return data[:8] == _OLE_MAGIC


def sniff_xlsx(data: bytes) -> bool:
    if data[:4] != b"PK\x03\x04":
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return "xl/workbook.xml" in z.namelist()
    except (zipfile.BadZipFile, ValueError, OSError):
        return False


def _guard_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        raise ValueError("not a valid .xlsx file") from None
    if len(infos) > MAX_ENTRIES:
        raise ValueError("spreadsheet has too many parts; refusing to open it")
    if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
        raise ValueError("spreadsheet expands to more than 200 MB; refusing to open it")


def _fmt_float(v: float) -> str:
    if v != v or v in (float("inf"), float("-inf")):
        return ""
    if float(v).is_integer() and abs(v) < 1e15:
        return str(int(v))
    s = format(round(v, 6), "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return "0" if s in ("-0", "") else s


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat()
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return _fmt_float(v)
    if isinstance(v, Decimal):
        return format(v.normalize(), "f")
    return str(v).strip()


_DATEISH = re.compile(r"^\d{1,4}[/.-]\d{1,2}[/.-]\d{1,4}$|^[A-Za-z]{3,9}\.? \d{1,2},? \d{4}$")
_NUMISH = re.compile(r"^\(?[-+]?\$?\s*[\d,]*\.?\d+\)?(?:\s*(?:CR|DR))?$", re.I)


def _is_data_cell(v) -> bool:
    if isinstance(v, (datetime, date)) or (isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)):
        return True
    if isinstance(v, str):
        s = v.strip()
        return bool(_DATEISH.match(s) or (_NUMISH.match(s) and re.search(r"\d", s)))
    return False


def _header_like(row: tuple) -> bool:
    texts = [c for c in row if isinstance(c, str) and c.strip() and not _is_data_cell(c)]
    return len(texts) >= 3


def _find_header(rows: list[tuple]) -> Optional[int]:
    for i, r in enumerate(rows):
        if not _header_like(r):
            continue
        for nxt in rows[i + 1:i + 6]:
            if any(_is_data_cell(c) for c in nxt):
                return i
    return None


def xlsx_to_csv_bytes(data: bytes, max_rows: int = 200000) -> tuple[bytes, dict]:
    """Convert the statement sheet of an .xlsx to UTF-8 CSV bytes.

    Returns (csv_bytes, info) with info = {"sheet", "sheets", "format": "XLSX", "header_row"}.
    Raises ValueError for legacy .xls, invalid files and zip bombs.
    """
    if is_legacy_xls(data):
        raise ValueError("Legacy .xls is not supported; re-save as .xlsx or CSV")
    if not sniff_xlsx(data):
        raise ValueError("not a valid .xlsx file")
    _guard_zip(data)
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"could not open the spreadsheet ({type(e).__name__})") from None
    try:
        sheets = wb.worksheets
        chosen = None
        header_row = None
        fallback = None
        for ws in sheets:
            head: list[tuple] = []
            for r in ws.iter_rows(max_row=60, values_only=True):
                head.append(tuple(r))
            if fallback is None and any(any(c is not None and str(c).strip() for c in r) for r in head):
                fallback = ws
            h = _find_header(head)
            if h is not None:
                chosen, header_row = ws, h
                break
        if chosen is None:
            chosen = fallback
        if chosen is None:
            raise ValueError("the spreadsheet has no data")
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        n = 0
        for r in chosen.iter_rows(values_only=True):
            n += 1
            if n > max_rows:
                break
            cells = [_cell_str(c) for c in r]
            while cells and cells[-1] == "":
                cells.pop()
            w.writerow(cells)
        info = {"sheet": chosen.title, "sheets": len(sheets), "format": "XLSX", "header_row": header_row}
        if n > max_rows:
            info["truncated"] = True
    finally:
        wb.close()
    return buf.getvalue().encode("utf-8"), info


def parse_xlsx(data: bytes, profile, account_id: str, import_id: str):
    """(txns, warnings, info) — the sheet is parsed by the CSV pipeline with `profile`."""
    from .csv_parser import parse_csv_detailed
    csv_bytes, info = xlsx_to_csv_bytes(data)
    txns, warnings = parse_csv_detailed(csv_bytes, profile, account_id, import_id)
    if info.get("truncated"):
        warnings = list(warnings) + ["spreadsheet was truncated at 200,000 rows"]
    return txns, warnings, info


def _csv_name(filename: str) -> str:
    name = filename or "upload.xlsx"
    new = re.sub(r"\.(?:xlsx|xlsm|xls)$", ".csv", name, flags=re.I)
    return new if new != name else name + ".csv"


def detect_xlsx(filename: str, data: bytes):
    """FormatProfile for the statement sheet (detected as if it were a CSV)."""
    from .detect import detect, read_rows
    csv_bytes, info = xlsx_to_csv_bytes(data)
    name = _csv_name(filename)
    prof = detect(name, csv_bytes)
    h = info.get("header_row")
    if h and (prof.mapping is None or (prof.header_row is not None and prof.header_row < h)):
        # A title/preamble row above the real header can fool the generic header
        # finder; retry on the sheet from the header we located, then shift back.
        lines = csv_bytes.decode("utf-8").split("\n")
        retry = detect(name, "\n".join(lines[h:]).encode("utf-8"))
        if retry.mapping is not None:
            if retry.header_row is not None and retry.header_row >= 0:
                retry.header_row += h
            prof = retry
    return prof
