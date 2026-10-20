# SPDX-License-Identifier: AGPL-3.0-only
from datetime import date, datetime
from decimal import Decimal

import openpyxl

from ledgerstitch.models import Direction, Lane, Transaction
from ledgerstitch.reports.excel import ALL_COLUMNS, SHEETS, export_excel
from ledgerstitch.reports.summary import DISCLAIMER

from test_summary import make_accounts, make_txns


def _build(tmp_path, txns=None, year=2025):
    p = export_excel(txns if txns is not None else make_txns(), make_accounts(),
                     tmp_path / "out" / "report.xlsx", "Sam's Studio LLC", year)
    assert p.exists()
    return openpyxl.load_workbook(p)


def _data_rows(ws, first=2):
    """Rows until the first blank column-A cell."""
    n = 0
    r = first
    while ws.cell(row=r, column=1).value not in (None, ""):
        n += 1
        r += 1
    return n


def test_sheets_and_headers(tmp_path):
    wb = _build(tmp_path)
    assert wb.sheetnames == SHEETS
    ws = wb["All Transactions"]
    assert [c.value for c in ws[1]] == ALL_COLUMNS
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.startswith("A1:L")


def test_row_counts(tmp_path):
    wb = _build(tmp_path)
    assert _data_rows(wb["All Transactions"]) == 9      # 2025, no duplicate
    assert _data_rows(wb["Business"]) == 4              # adobe, refund, meal, fee
    assert _data_rows(wb["Personal"]) == 1
    assert _data_rows(wb["Needs Review"]) == 1
    assert _data_rows(wb["Excluded"]) == 4              # 2 transfer legs + income + duplicate
    reasons = [wb["Excluded"].cell(row=r, column=1).value for r in range(2, 6)]
    assert reasons.count("Duplicate") == 1
    assert "Income (not an expense)" in reasons
    assert reasons.count("Transfer / card payment") == 2


def test_all_transactions_values(tmp_path):
    wb = _build(tmp_path)
    ws = wb["All Transactions"]
    first = [c.value for c in ws[2]]
    assert isinstance(first[0], datetime) and first[0].date() == date(2025, 1, 5)
    assert first[1] == "Chase Freedom ••4012"
    assert first[7] == -100.0                           # signed: out negative
    assert first[9] == "Rule"                           # audit trail source
    assert "#,##0.00" in ws.cell(row=2, column=8).number_format
    amounts = [ws.cell(row=r, column=8).value for r in range(2, 11)]
    assert round(sum(amounts), 2) == round(-100 + 20 - 60 - 50 - 30 - 500 + 500 + 2000 - 15, 2)
    total = ws.cell(row=12, column=8).value
    assert total == "=SUBTOTAL(9,H2:H10)"


def test_lane_tab_totals_are_subtotal_formulas(tmp_path):
    wb = _build(tmp_path)
    ws = wb["Business"]
    vals = [ws.cell(row=r, column=7).value for r in range(2, 6)]
    assert sum(vals) == 155.0                           # refund nets as negative
    assert ws.cell(row=7, column=7).value == "=SUBTOTAL(9,G2:G5)"


def test_summary_static_values_and_links(tmp_path):
    wb = _build(tmp_path)
    ws = wb["Summary"]
    assert ws["A1"].value == "Sam's Studio LLC"
    assert ws["B2"].value == 2025
    assert ws["A4"].value == DISCLAIMER
    labels = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=2).value,
                                               ws.cell(row=r, column=3).value)
              for r in range(7, 14)}
    assert labels["Business expenses"][0] == 155.0
    assert labels["Business expenses"][1] == "='Business'!G7"
    assert labels["Personal spending"][0] == 50.0
    assert labels["Needs review (unresolved)"][0] == 30.0
    assert labels["Income (not an expense)"][0] == 2000.0
    assert labels["Transfers out (excluded)"][0] == 500.0
    assert labels["Total expenses (business + personal + review)"][0] == 235.0
    # category hyperlinks into All Transactions
    links = {}
    for row in ws.iter_rows(min_row=1):
        c = row[0]
        if c.hyperlink is not None:
            links[c.value] = c.hyperlink.location
    assert links["Software & Subscriptions"] == "'All Transactions'!A2"
    assert links["Groceries"] == "'All Transactions'!A5"


def test_schedule_c_tab(tmp_path):
    wb = _build(tmp_path)
    ws = wb["Schedule C"]
    rows = {ws.cell(row=r, column=1).value: [ws.cell(row=r, column=c).value for c in range(2, 7)]
            for r in range(6, 9)}
    assert set(rows) == {"L10", "L18", "L24b"}
    assert rows["L24b"][1:4] == [60.0, 0.5, 30.0]
    assert rows["L18"][1:4] == [80.0, 1.0, 80.0]
    texts = [ws.cell(row=r, column=1).value for r in range(1, 20)]
    assert DISCLAIMER in texts
    assert any(t and "50%" in t for t in texts)


def test_by_account(tmp_path):
    wb = _build(tmp_path)
    ws = wb["By Account"]
    rows = {ws.cell(row=r, column=1).value: [ws.cell(row=r, column=c).value for c in range(4, 11)]
            for r in range(2, 4)}
    assert rows["Chase Freedom ••4012"][0] == 140.0
    assert rows["Chase Checking ••1111"][3] == 500.0   # transfers out
    assert rows["Chase Checking ••1111"][5] == 2000.0  # income
    assert ws.cell(row=4, column=1).value == "All accounts"
    assert ws.cell(row=4, column=4).value == "=SUBTOTAL(9,D2:D3)"


def test_formula_injection_is_text(tmp_path):
    t = Transaction("chk-1111", date(2025, 5, 1), "=HYPERLINK(\"http://x\")", Decimal("1"),
                    Direction.OUT, lane=Lane.PERSONAL, category="Shopping")
    wb = _build(tmp_path, [t])
    c = wb["All Transactions"]["C2"]
    assert c.data_type == "s"


def test_empty_workbook(tmp_path):
    wb = _build(tmp_path, [])
    assert wb.sheetnames == SHEETS
