# SPDX-License-Identifier: AGPL-3.0-only
"""1099 import, holdings snapshot and investment signals, on SYNTHETIC statements
(tests/fixtures/make_1099s.py: Example Brokerage, Sample Securities, Specimen Markets; all invented)."""
from __future__ import annotations

import csv
import io
import json
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import pytest

from ledgerstitch.cli import main as cli_main
from ledgerstitch.investments import (DISCLAIMER, LABEL, PLUS_FEATURES, analyze_investments, parse_1099,
                                      parse_1099_text, parse_lots_csv, parse_positions_csv, summarize_holdings)
from ledgerstitch.investments.forms1099 import category_for, mask_tax_ids
from ledgerstitch.investments.holdings import asset_class
from ledgerstitch.store import to_jsonable

FX = Path(__file__).parent / "fixtures"
SECRETS = ["900-45-6789", "98-7654321", "900123456", "900-234567", "900345678", "900456789", "Z90012345"]
ALL = ["1099_example_brokerage.pdf", "1099_sample_securities.pdf", "1099_specimen_markets.txt",
       "1099r_sample_trust.txt", "1099nec_example_client.txt", "1099b_lots_example.csv",
       "positions_example.csv", "positions_sample.csv"]


def load(*names):
    return [(n, (FX / n).read_bytes()) for n in names]


def stmt(name):
    return parse_1099(name, (FX / name).read_bytes())


def household(**kw):
    return analyze_investments(load(*ALL), **kw)


def by_id(signals, sid):
    return [s for s in signals if s.id == sid]


# --------------------------------------------------------------------------- fixtures are current
def test_fixtures_match_generator():
    from fixtures.make_1099s import FIXTURES
    for name, make in FIXTURES.items():
        assert (FX / name).read_bytes() == make(), f"{name} is stale: run tests/fixtures/make_1099s.py"


# --------------------------------------------------------------------------- 1099 layouts
def test_fidelity_style_pdf():
    s = stmt("1099_example_brokerage.pdf")
    assert s.payer == "Example Brokerage LLC" and s.tax_year == 2025
    assert s.account == "XXXXX3456" and s.recipient_tin == "XXX-XX-6789"
    div = s.forms["1099-DIV"]
    assert div["1a"] == D("12032.46") and div["1b"] == D("9410.22") and div["2a"] == D("1204.18")
    assert div["5"] == D("800.00") and div["7"] == D("245.30") and div["12"] == D("310.00") and div["13"] == D("22.00")
    assert s.text_boxes["1099-DIV"]["8"] == "Various"
    assert s.forms["1099-INT"] == {"1": D("412.10"), "3": D("150.00"), "4": D("0.00")}
    b = {g.category: g for g in s.b_summary}
    assert set(b) == {"A", "B", "D"}
    assert (b["A"].proceeds, b["A"].cost_basis, b["A"].wash_sale_disallowed, b["A"].gain) == \
        (D("8050.00"), D("7400.00"), D("200.00"), D("850.00"))
    assert b["B"].proceeds == D("3000.00") and b["B"].cost_basis is None and b["B"].gain is None
    assert b["D"].gain == D("3400.00") and b["D"].term == "long"
    assert any("no cost basis" in w for w in s.warnings)


def test_schwab_style_pdf():
    s = stmt("1099_sample_securities.pdf")
    assert s.payer == "Sample Securities Inc." and s.tax_year == 2025
    assert s.forms["1099-DIV"] == {"1a": D("747.04"), "1b": D("600.00"), "2a": D("120.00"), "3": D("50.00"),
                                   "7": D("12.50")}
    assert s.forms["1099-INT"] == {"1": D("88.20")} and s.forms["1099-MISC"] == {"3": D("25.00")}
    b = {g.category: (g.proceeds, g.cost_basis, g.gain) for g in s.b_summary}
    assert b == {"A": (D("20000.00"), D("26000.00"), D("-6000.00")), "D": (D("5000.00"), D("4000.00"), D("1000.00")),
                 "E": (D("2500.00"), D("2000.00"), D("500.00"))}
    assert s.warnings == []


def test_robinhood_style_text():
    s = stmt("1099_specimen_markets.txt")
    assert s.payer == "Specimen Markets LLC" and s.tax_year == 2025 and s.account == "XXXXX5678"
    assert s.forms["1099-DIV"] == {"1a": D("210.00"), "1b": D("150.00"), "5": D("40.00"), "4": D("0.00")}
    [a] = s.b_summary                         # the "Total Short-term" row is not a category
    assert (a.category, a.proceeds, a.cost_basis, a.market_discount, a.wash_sale_disallowed, a.gain) == \
        ("A", D("9500.00"), D("10000.00"), D("0.00"), D("300.00"), D("-200.00"))


def test_1099_r_and_nec():
    r = stmt("1099r_sample_trust.txt")
    assert r.payer == "Sample Retirement Trust" and r.forms["1099-R"] == {"1": D("5000.00"), "2a": D("5000.00"),
                                                                          "4": D("1000.00")}
    assert r.text_boxes["1099-R"]["7"] == "1"
    n = stmt("1099nec_example_client.txt")
    assert n.payer == "Example Client Inc" and n.forms["1099-NEC"]["1"] == D("2500.00")


def test_label_matching_without_box_numbers():
    s = parse_1099_text("Example Brokerage LLC\nForm 1099-DIV\nTotal ordinary dividends      100.00\n"
                        "Qualified dividends   80.00\nForeign tax paid   3.00\nForm 1099-INT\n"
                        "Interest income   9.99\nTax-exempt interest   1.50\n")
    assert s.forms == {"1099-DIV": {"1a": D("100.00"), "1b": D("80.00"), "7": D("3.00")},
                       "1099-INT": {"1": D("9.99"), "8": D("1.50")}}


def test_1099b_without_headers_by_count_and_printed_net_check():
    s = parse_1099_text("Form 1099-B\nShort-term covered (Box A)   1,000.00   900.00   0.00   50.00   200.00\n"
                        "Long-term noncovered   500.00   400.00   100.00\n")
    a, e = s.b_summary
    assert (a.category, a.proceeds, a.cost_basis, a.wash_sale_disallowed, a.gain) == \
        ("A", D("1000.00"), D("900.00"), D("50.00"), D("200.00"))
    assert any("printed net 200.00 differs" in w for w in s.warnings)     # 1000 - 900 + 50 = 150
    assert (e.category, e.gain) == ("E", D("100.00"))


@pytest.mark.parametrize("term,covered,reported,box", [
    ("short", True, True, "A"), ("short", False, True, "B"), ("short", None, False, "C"),
    ("long", True, True, "D"), ("long", False, True, "E"), ("long", None, False, "F"), ("unknown", True, True, "X"),
])
def test_category_for(term, covered, reported, box):
    assert category_for(term, covered, reported) == box


def test_no_text_layer_and_garbage():
    from ledgerstitch.bench.pdfgen import PDF
    s = parse_1099("scan.pdf", PDF().to_bytes())
    assert any("no text layer" in w for w in s.warnings)
    s = parse_1099("broken.pdf", b"%PDF-1.4 garbage")
    assert any("could not read PDF" in w for w in s.warnings)
    assert any("no 1099 boxes" in w for w in parse_1099_text("hello").warnings)


def test_ein_is_masked():
    assert mask_tax_ids("Payer federal ID 98-7654321") == "Payer federal ID XX-XXX4321"
    assert mask_tax_ids("Payer 98-7654321 and 2025-01-31") == "Payer XX-XXX4321 and 2025-01-31"


# --------------------------------------------------------------------------- 1099-B lots
def test_lots_csv():
    s = parse_lots_csv("1099b_lots_example.csv", (FX / "1099b_lots_example.csv").read_bytes())
    assert s.kind == "1099-B lots" and s.payer == "Example Brokerage LLC" and s.tax_year == 2025
    assert s.account == "XXXXX3456" and len(s.lots) == 5            # the Total row is skipped
    lg = next(lt for lt in s.lots if lt.symbol == "LGCY")
    assert lg.date_acquired is None and lg.cost_basis is None and lg.category == "B" and lg.gain is None
    sm = next(lt for lt in s.lots if lt.symbol == "SMPM")
    assert sm.gain == D("-400.00") and sm.wash_sale_disallowed == D("200.00") and sm.category == "A"
    assert {g.category: g.proceeds for g in s.b_summary} == {"A": D("8050.00"), "B": D("3000.00"), "D": D("14000.00")}


def test_lots_term_from_dates_and_box_column():
    data = ("Security,Date Acquired,Date Sold,Proceeds,Cost Basis,Box\n"
            "AAA,01/02/2024,01/02/2025,100.00,90.00,\n"          # exactly one year: short-term
            "BBB,01/02/2024,01/03/2025,100.00,90.00,\n"
            "CCC,,06/01/2025,50.00,60.00,(E)\n").encode()
    s = parse_lots_csv("lots.csv", data)
    assert [(lt.term, lt.category) for lt in s.lots] == [("short", "A"), ("long", "D"), ("long", "E")]


def test_lots_matching_a_summary_are_counted_once():
    r = analyze_investments(load("1099_example_brokerage.pdf", "1099b_lots_example.csv"))
    assert {g.category: g.proceeds for g in r.capital_gains} == {"A": D("8050.00"), "B": D("3000.00"),
                                                                 "D": D("14000.00")}
    assert any("counted once" in w for w in r.warnings)
    alone = analyze_investments(load("1099b_lots_example.csv"))
    assert alone.short_term_net == D("850.00") and alone.long_term_net == D("3400.00")


def test_lots_not_matching_are_added_with_a_warning():
    r = analyze_investments(load("1099_sample_securities.pdf", "1099b_lots_example.csv"))
    assert {g.category: g.proceeds for g in r.capital_gains}["A"] == D("28050.00")
    assert any("don't match any 1099-B summary" in w for w in r.warnings)


# --------------------------------------------------------------------------- positions
def test_fidelity_positions():
    hs, as_of, w = parse_positions_csv("positions_example.csv", (FX / "positions_example.csv").read_bytes())
    assert as_of is None and w == []
    assert [h.symbol for h in hs] == ["SPAXX", "EXTC", "BMIX", "EPPX", "SBIO", "VXFDX"]   # no Pending Activity
    assert all(h.account == "ZXXXX2345" for h in hs)
    cls = {h.symbol: h.asset_class for h in hs}
    assert cls == {"SPAXX": "cash", "EXTC": "equity", "BMIX": "etf", "EPPX": "equity", "SBIO": "equity",
                   "VXFDX": "fund"}
    sbio = next(h for h in hs if h.symbol == "SBIO")
    assert (sbio.quantity, sbio.market_value, sbio.cost_basis, sbio.unrealized_gain) == \
        (D("200"), D("4000.00"), D("7000.00"), D("-3000.00"))


def test_schwab_positions():
    hs, as_of, w = parse_positions_csv("positions_sample.csv", (FX / "positions_sample.csv").read_bytes())
    assert as_of == date(2025, 10, 24) and w == []
    assert [h.asset_class for h in hs] == ["equity", "option", "bond", "cash"]      # Account Total skipped
    assert all(h.account == "Individual ...898" for h in hs)
    opt = hs[1]
    assert opt.market_value == D("100.00") and opt.unrealized_gain == D("-300.00")
    assert hs[3].market_value == D("1400.00") and hs[3].quantity is None


def test_vanguard_style_positions_value_from_price():
    data = ("Account Number,Investment Name,Symbol,Shares,Share Price,Total Value\n"
            "900111222,SAMPLE 500 INDEX ADMIRAL,VSMPX,10,400.00,\n"
            "900111222,SAMPLE FEDERAL MONEY MARKET,VMFXX,1000,1.00,1000.00\n").encode()
    hs, _, _ = parse_positions_csv("vg.csv", data)
    assert [(h.symbol, h.market_value, h.asset_class, h.account) for h in hs] == [
        ("VSMPX", D("4000.00"), "fund", "XXXXX1222"), ("VMFXX", D("1000.00"), "cash", "XXXXX1222")]
    assert hs[0].cost_basis is None and hs[0].unrealized_gain is None


@pytest.mark.parametrize("sym,desc,typ,expected", [
    ("AAPL250117C00150000", "", "", "option"),
    ("XYZ", "PUT XYZ $10 EXP 01/16/26", "", "option"),
    ("FDRXX", "", "", "cash"),
    ("BTC", "Bitcoin", "", "crypto"),
    ("ETHA", "ISHARES ETHEREUM TRUST ETF", "", "etf"),
    ("912828XYZ", "US TREASURY NOTE 4.25% 01/31/2027", "", "bond"),
    ("ABCDX", "SAMPLE GROWTH FUND CLASS A", "", "fund"),
    ("ABC", "ABC CORP", "Equity", "equity"),
    ("", "", "", "other"),
])
def test_asset_class(sym, desc, typ, expected):
    assert asset_class(sym, desc, typ) == expected


def test_holdings_summary():
    r = household()
    s = r.holdings_summary
    assert s.positions == 10 and s.market_value == D("55000.00") and s.as_of == date(2025, 10, 24)
    assert s.unrealized_losses == D("-4800.00") and s.loss_positions == 4 and s.basis_unknown == 0
    assert s.unrealized_gain == D("4700.00")
    assert [(a.name, a.pct) for a in s.concentrated] == [("EXTC", D("32.7"))]
    assert sum(a.value for a in s.by_class) == s.market_value
    assert {a.name for a in s.by_account} == {"ZXXXX2345", "Individual ...898"}
    assert "SPAXX" not in {a.name for a in s.top}
    assert summarize_holdings([]).market_value == 0


# --------------------------------------------------------------------------- household report + signals
def test_household_totals():
    r = household(filing_status="mfj")
    assert r.year == 2025 and r.label == LABEL and r.disclaimer == DISCLAIMER and r.plus_features == PLUS_FEATURES
    inc = {(t.form, t.box): t for t in r.income}
    assert inc[("1099-DIV", "1a")].amount == D("12989.50")
    assert inc[("1099-DIV", "1a")].payers == ["Example Brokerage LLC", "Sample Securities Inc.", "Specimen Markets LLC"]
    assert inc[("1099-DIV", "1b")].destination == "Form 1040 line 3a"
    assert inc[("1099-INT", "1")].amount == D("500.30") and inc[("1099-R", "4")].amount == D("1000.00")
    assert [t.form for t in r.income] == sorted((t.form for t in r.income),
                                                key=["1099-DIV", "1099-INT", "1099-MISC", "1099-NEC",
                                                     "1099-R"].index)
    gains = {g.category: g for g in r.capital_gains}
    assert gains["A"].proceeds == D("37550.00") and gains["A"].gain == D("-5350.00")
    assert gains["B"].cost_basis is None and "Schedule D line 2" in gains["B"].destination
    assert r.short_term_net == D("-5350.00") and r.long_term_net == D("4900.00")


def test_household_signals():
    r = household(filing_status="mfj")
    ids = {s.id for s in r.signals}
    assert {"unrealized_losses", "basis_missing", "noncovered_basis", "wash_sale_disallowed", "foreign_tax_paid",
            "section_199a_dividends", "exempt_interest", "nondividend_distributions", "partnership_k1_likely",
            "early_distribution"} <= ids
    assert "no_withholding" not in ids            # the 1099-R withheld $1,000
    assert "short_term_gains" not in ids          # net short-term is a loss
    [loss] = by_id(r.signals, "unrealized_losses")
    assert loss.amount == D("4800.00") and loss.severity == "money" and "LedgerStitch Plus" in loss.explanation
    [b] = by_id(r.signals, "basis_missing")
    assert b.amount == D("3000.00") and b.severity == "warning"
    [w] = by_id(r.signals, "wash_sale_disallowed")
    assert w.amount == D("500.00")
    [f] = by_id(r.signals, "foreign_tax_paid")
    assert f.amount == D("257.80") and "without Form 1116" in f.explanation
    [q] = by_id(r.signals, "section_199a_dividends")
    assert q.amount == D("168.00")                # 20% of $840
    [k1] = by_id(r.signals, "partnership_k1_likely")
    assert "EPPX" in k1.title
    [e] = by_id(r.signals, "early_distribution")
    assert e.amount == D("5000.00")
    for s in r.signals:
        assert s.citation and s.explanation


def test_no_withholding_and_short_term_signals():
    r = analyze_investments(load("1099_example_brokerage.pdf", "1099nec_example_client.txt"))
    [nw] = by_id(r.signals, "no_withholding")
    assert nw.amount == D("12032.46") + D("1204.18") + D("412.10") + D("2500.00") + D("850.00") + D("3400.00")
    assert "short_term_gains" not in {s.id for s in r.signals}      # long-term is larger
    st = analyze_investments([("b.txt", b"Form 1099-B\nShort-term covered (Box A)   9,000.00   1,000.00   8,000.00\n")])
    assert by_id(st.signals, "short_term_gains")[0].amount == D("8000.00")


def test_capital_loss_limit_and_foreign_tax_1116():
    text = ("Form 1099-DIV\n7 Foreign tax paid   450.00\nForm 1099-B\n"
            "Short-term covered (Box A)   1,000.00   11,000.00   -10,000.00\n").encode()
    r = analyze_investments([("x.txt", text)], filing_status="single")
    [c] = by_id(r.signals, "capital_loss_limit")
    assert c.amount == D("7000.00")
    assert "Form 1116 is required" in by_id(r.signals, "foreign_tax_paid")[0].explanation
    mfs = analyze_investments([("x.txt", text)], filing_status="mfs")
    assert by_id(mfs.signals, "capital_loss_limit")[0].amount == D("8500.00")


def test_year_filter_and_validation():
    old = b"Example Brokerage LLC\nTax Year 2024\nForm 1099-INT\n1 Interest income   5.00\n"
    r = analyze_investments(load("1099_specimen_markets.txt") + [("old.txt", old)])
    assert r.year == 2025 and any("tax year 2024, not 2025; skipped" in w for w in r.warnings)
    assert analyze_investments([("old.txt", old)], year=2024).income[0].amount == D("5.00")
    with pytest.raises(ValueError):
        analyze_investments([], filing_status="joint")
    assert "no files given" in analyze_investments([]).warnings


# --------------------------------------------------------------------------- no identifiers anywhere
def test_no_identifiers_in_outputs(tmp_path):
    from openpyxl import load_workbook

    from ledgerstitch.reports.investments_export import export_investments_csv, export_investments_xlsx
    r = household()
    blob = json.dumps(to_jsonable(r))
    csv_text = export_investments_csv(r, tmp_path / "i.csv").read_text()
    wb = load_workbook(export_investments_xlsx(r, tmp_path / "i.xlsx"))
    cells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    for secret in SECRETS:
        assert secret not in blob and secret not in csv_text and secret not in cells, secret


def test_csv_export_layout(tmp_path):
    from ledgerstitch.reports.investments_export import INVESTMENT_COLUMNS, export_investments_csv
    rows = list(csv.reader(io.StringIO(export_investments_csv(household(), tmp_path / "i.csv").read_text())))
    assert rows[0] == [LABEL] and rows[2] == INVESTMENT_COLUMNS
    sections = {r[0] for r in rows[3:]}
    assert {"1099 income", "Capital gains", "Holding", "Portfolio", "Signal", "Statement", "Warning"} <= sections
    div = next(r for r in rows[3:] if r[2].startswith("1099-DIV box 1a"))
    assert div[3] == "12989.50" and "line 3b" in div[4]


def test_xlsx_export_tabs(tmp_path):
    from openpyxl import load_workbook

    from ledgerstitch.reports.investments_export import export_investments_xlsx
    wb = load_workbook(export_investments_xlsx(household(), tmp_path / "i.xlsx"))
    assert wb.sheetnames == ["1099 income", "Capital gains", "Holdings", "Signals", "Statements"]
    assert wb["1099 income"]["A1"].value == LABEL


# --------------------------------------------------------------------------- CLI
def test_cli_investments(capsys, tmp_path):
    out = tmp_path / "inv.xlsx"
    rc = cli_main(["investments", *(str(FX / n) for n in ALL), "--filing-status", "mfj", "--out", str(out)])
    text = capsys.readouterr().out
    assert rc == 0 and out.exists()
    assert LABEL in text and "1099-DIV  box 1a" in text and "concentrated: EXTC" in text
    assert "In LedgerStitch Plus:" in text
    for secret in SECRETS:
        assert secret not in text


def test_cli_investments_json(capsys):
    assert cli_main(["investments", str(FX / "1099_specimen_markets.txt"), "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["year"] == 2025 and data["capital_gains"][0]["gain"] == "-200.00"


def test_cli_investments_missing_file(tmp_path):
    assert cli_main(["investments", str(tmp_path / "nope.pdf")]) == 1
