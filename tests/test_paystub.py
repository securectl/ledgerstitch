# SPDX-License-Identifier: AGPL-3.0-only
"""Pay stub import, projection, estimated W-2 and household signals, on SYNTHETIC stubs
(tests/fixtures/make_paystubs.py: Acme Corp, Initech Inc, Globex LLC, Jane Doe; all invented)."""
from __future__ import annotations

import csv
import io
import json
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import pytest

from ledgerstitch.cli import main as cli_main
from ledgerstitch.payroll import analyze_paystubs, estimate_w2, parse_paystub, parse_paystub_text
from ledgerstitch.payroll.limits import LIMITS, limits_for
from ledgerstitch.payroll.paystub import classify_deduction, parse_amounts
from ledgerstitch.payroll.signals import household_signals
from ledgerstitch.payroll.w2 import ESTIMATE_LABEL, detect_frequency, periods_remaining
from ledgerstitch.store import to_jsonable

FX = Path(__file__).parent / "fixtures"
SECRETS = ["900-45-6789", "987-65-4321", "000123456789", "9876543210", "00987654", "01234567"]


def stub(name: str, owner: str = "you"):
    return parse_paystub(name, (FX / name).read_bytes(), owner)


def household(**kw):
    files = [("paystub_rippling.pdf", (FX / "paystub_rippling.pdf").read_bytes(), "you"),
             ("paystub_generic.txt", (FX / "paystub_generic.txt").read_bytes(), "you"),
             ("paystub_adp.pdf", (FX / "paystub_adp.pdf").read_bytes(), "spouse")]
    return analyze_paystubs(files, **kw)


def by_id(signals, sid, owner="*"):
    return [s for s in signals if s.id == sid and (owner == "*" or s.owner == owner)]


# --------------------------------------------------------------------------- fixtures are current
def test_fixtures_match_generator():
    from fixtures.make_paystubs import FIXTURES
    for name, make in FIXTURES.items():
        assert (FX / name).read_bytes() == make(), f"{name} is stale: run tests/fixtures/make_paystubs.py"


# --------------------------------------------------------------------------- amounts
@pytest.mark.parametrize("line,values,stars", [
    ("Regular   85.5411   80.00      6 843 29     47 500 02", ["80.00", "6843.29", "47500.02"], [0, 0, 0]),
    ("401K Contrib            -487 36*      3 365 88", ["-487.36", "3365.88"], [1, 0]),
    ("Gross Pay      $6 843 29     73 967 62", ["6843.29", "73967.62"], [0, 0]),
    ("Fed      1 234 567 89", ["1234567.89"], [0]),
    ("Salary      5,000.00   55,000.00", ["5000.00", "55000.00"], [0, 0]),
    ("Refund      (12.00)   15.00-", ["-12.00", "-15.00"], [0, 0]),
    ("Section 125 Medical   120.00", ["120.00"], [0]),     # "125" is part of the label, not money
])
def test_parse_amounts(line, values, stars):
    amts, _ = parse_amounts(line)
    assert [str(a.value) for a in amts] == values
    assert [int(a.star) for a in amts] == stars


@pytest.mark.parametrize("label,kind,fsa", [
    ("401(k) Pre-Tax", "pretax_401k", None), ("401K Contrib", "pretax_401k", None),
    ("Roth 401(k)", "roth_401k", None), ("403(b)", "pretax_403b", None), ("Roth 403b", "roth_403b", None),
    ("HSA", "hsa", None), ("Health Savings Acct", "hsa", None),
    ("Health FSA", "health_fsa", "unknown"), ("Limited Purpose FSA", "health_fsa", "limited"),
    ("Dep Care FSA", "dcap", None), ("Dependent Care", "dcap", None),
    ("Medical", "sec125_medical", None), ("Dental", "sec125_dental", None), ("Vision", "sec125_vision", None),
    ("Life Ins", "post_tax", None), ("Critical Illness", "post_tax", None), ("Accident", "post_tax", None),
    ("Hospital Indemnity", "post_tax", None), ("401k Loan", "post_tax", None),
    ("Expense Reimbursement", "reimbursement", None), ("Wellness Plan", "other", None),
])
def test_classify_deduction(label, kind, fsa):
    assert classify_deduction(label) == (kind, fsa)


# --------------------------------------------------------------------------- layouts
def test_rippling_pdf():
    st = stub("paystub_rippling.pdf")
    assert st.layout == "rippling" and st.employer == "Acme Corp"
    assert (st.period_start, st.period_end, st.pay_date) == (date(2025, 6, 1), date(2025, 6, 15), date(2025, 6, 15))
    assert st.filing_status == "single"
    assert (st.gross_current, st.gross_ytd) == (D("5004.50"), D("55049.50"))
    assert (st.net_current, st.net_ytd) == (D("3040.49"), D("33445.39"))
    assert [(e.label, e.kind) for e in st.earnings] == [("Salary", "wages"), ("Group Term Life", "imputed_gtl")]
    assert (st.federal_income_tax.current, st.federal_income_tax.ytd) == (D("450.00"), D("4950.00"))
    assert (st.social_security.current, st.social_security.ytd) == (D("280.83"), D("3089.13"))
    assert (st.medicare.current, st.medicare.ytd) == (D("65.68"), D("722.48"))
    assert [t.label for t in st.other_taxes] == ["IL State Income Tax"]
    d = {x.label: x for x in st.deductions}
    assert set(d) == {"401(k) Pre-Tax", "Medical", "Dental", "Vision", "HSA", "Health FSA", "Life Insurance"}
    k = d["401(k) Pre-Tax"]
    assert (k.current, k.ytd, k.employer_current, k.employer_ytd) == (D("500.00"), D("5500.00"), D("250.00"),
                                                                      D("2750.00"))
    assert k.pretax is True and k.fica_exempt is False and k.pretax_source == "stub"
    assert d["Medical"].employer_current == D("600.00") and d["Medical"].fica_exempt
    assert d["Health FSA"].fsa_type == "unknown"
    assert d["Life Insurance"].pretax is False
    assert st.checks[0]["check"] == "social_security_wages" and st.checks[0]["ok"]
    assert st.warnings == []


def test_adp_pdf_and_text_parse_the_same():
    a, b = stub("paystub_adp.pdf"), stub("paystub_adp.txt")
    ja, jb = to_jsonable(a), to_jsonable(b)
    ja.pop("source"), jb.pop("source")
    assert ja == jb


def test_adp_space_separated_amounts_and_stars():
    st = stub("paystub_adp.txt", "spouse")
    assert st.layout == "adp" and st.employer == "INITECH INC" and st.owner == "spouse"
    assert (st.period_start, st.period_end, st.pay_date) == (date(2025, 3, 1), date(2025, 3, 14), date(2025, 3, 21))
    assert st.filing_status == "married filing jointly"
    assert (st.gross_current, st.gross_ytd) == (D("6843.29"), D("73967.62"))
    reg, bonus = st.earnings
    assert (reg.current, reg.ytd, reg.hours) == (D("6843.29"), D("47500.02"), D("80.00"))
    assert (bonus.current, bonus.ytd) == (D("0"), D("26467.60"))          # YTD-only line
    assert st.federal_income_tax.current == D("1012.44")                    # "-1 012 44" → positive tax
    assert st.social_security.current == D("401.05") and st.medicare.ytd == D("728.66")
    assert st.federal_taxable_current == D("5981.13")
    d = {x.label: x for x in st.deductions}
    assert d["401K Contrib"].current == D("487.36") and d["401K Contrib"].pretax_source == "stub"
    assert d["Dep Care FSA"].kind == "dcap" and d["Dep Care FSA"].pretax
    assert d["Life Ins"].pretax is False
    assert st.net_current == D("3986.74")
    assert "Checking1" not in {x.label for x in st.deductions}             # deposit section ignored


def test_generic_text():
    st = stub("paystub_generic.txt")
    assert st.layout == "generic" and st.employer == "Globex LLC"
    assert (st.period_start, st.period_end) == (date(2025, 6, 1), date(2025, 6, 30))
    assert (st.gross_current, st.gross_ytd) == (D("9000.00"), D("54000.00"))
    assert st.social_security.ytd == D("3348.00") and st.medicare.current == D("130.50")
    (roth,) = st.deductions
    assert roth.kind == "roth_401k" and roth.pretax is False
    assert st.earnings == []


def test_generic_without_headers_label_current_ytd():
    st = parse_paystub_text("Company: Acme Corp\nCheck Date: 01/31/2026\nPay period 01/01/2026 - 01/31/2026\n"
                            "Salary 4,000.00 4,000.00\nGross Pay 4,000.00 4,000.00\nFederal Withholding 400.00 400.00\n"
                            "OASDI 248.00 248.00\nMedicare 58.00 58.00\nNet Pay 3,294.00 3,294.00\n")
    assert st.employer == "Acme Corp" and st.pay_date == date(2026, 1, 31)
    assert st.federal_income_tax.current == D("400.00") and st.social_security.current == D("248.00")
    assert [e.label for e in st.earnings] == ["Salary"]


def test_unreadable_inputs_warn():
    assert parse_paystub("x.pdf", b"%PDF-1.4 garbage").warnings[0].startswith("could not read PDF")
    st = parse_paystub("x.txt", b"hello world\n")
    assert any("no earnings" in w for w in st.warnings)


# --------------------------------------------------------------------------- pre-tax inference
def test_pretax_inferred_from_fica_math():
    st = stub("paystub_adp.txt")
    w = next(x for x in st.deductions if x.label == "Wellness Plan")
    assert w.kind == "other" and w.pretax is True and w.fica_exempt is True
    assert w.pretax_source == "inferred (FICA)"
    chk = next(c for c in st.checks if c["check"] == "social_security_wages")
    assert chk["ok"] and chk["inferred"] == ["Wellness Plan (pre-tax for FICA)"]
    fed = next(c for c in st.checks if c["check"] == "federal_taxable_wages")
    assert fed["ok"]


def test_fica_math_flips_assumed_pretax_medical_to_post_tax():
    # SS = 6.2% of the full gross: the "Medical" line is post-tax here.
    st = parse_paystub_text("Acme Corp\nPay period: 01/01/2025 - 01/15/2025\n"
                            "Description    Current     YTD\nGross Pay    2,000.00   2,000.00\n"
                            "Social Security    124.00   124.00\nMedicare    29.00   29.00\n"
                            "Medical    50.00   50.00\n")
    med = st.deductions[0]
    assert med.kind == "sec125_medical" and med.fica_exempt is False and med.pretax_source == "inferred (FICA)"


def test_fica_mismatch_is_a_warning_not_a_guess():
    st = parse_paystub_text("Acme Corp\nPay period: 01/01/2025 - 01/15/2025\n"
                            "Description    Current     YTD\nGross Pay    2,000.00   2,000.00\n"
                            "Social Security    100.00   100.00\nMedical    50.00   50.00\n")
    assert any("Social Security tax implies" in w for w in st.warnings)
    assert st.deductions[0].pretax_source == "assumed"


# --------------------------------------------------------------------------- frequency / projection
@pytest.mark.parametrize("start,end,freq", [
    (date(2025, 6, 1), date(2025, 6, 15), "semimonthly"), (date(2025, 2, 16), date(2025, 2, 28), "semimonthly"),
    (date(2025, 3, 1), date(2025, 3, 14), "biweekly"), (date(2025, 3, 3), date(2025, 3, 9), "weekly"),
    (date(2025, 6, 1), date(2025, 6, 30), "monthly"), (date(2025, 6, 1), date(2025, 6, 3), None), (None, None, None),
])
def test_detect_frequency(start, end, freq):
    assert detect_frequency(start, end) == freq


@pytest.mark.parametrize("pay,freq,n", [
    (date(2025, 6, 15), "semimonthly", 13), (date(2025, 6, 30), "semimonthly", 12),
    (date(2025, 3, 21), "biweekly", 20), (date(2025, 12, 26), "weekly", 0), (date(2025, 12, 19), "weekly", 1),
    (date(2025, 6, 30), "monthly", 6), (date(2025, 12, 31), "monthly", 0),
])
def test_periods_remaining(pay, freq, n):
    assert periods_remaining(pay, freq) == n


def test_projection_and_final():
    st = stub("paystub_rippling.pdf")
    proj = estimate_w2(st)
    assert proj.frequency == "semimonthly" and proj.periods_remaining == 13
    assert proj.gross == D("120108.00")                      # 55,049.50 + 13 × 5,004.50
    fin = estimate_w2(st, final=True)
    assert fin.periods_remaining == 0 and fin.gross == D("55049.50") and fin.boxes["2"] == D("4950.00")
    forced = estimate_w2(st, frequency="monthly")
    assert forced.periods_remaining == 6


def test_unknown_frequency_falls_back_to_final():
    st = parse_paystub_text("Acme Corp\nGross Pay  1,000.00  5,000.00\n")
    est = estimate_w2(st, year=2025)
    assert est.final and est.gross == D("5000.00") and "treated as final" in est.notes[0]


# --------------------------------------------------------------------------- estimated W-2
def test_w2_boxes_rippling():
    w = estimate_w2(stub("paystub_rippling.pdf"))
    assert w.label == ESTIMATE_LABEL == "ESTIMATE — not issued by employer"
    assert w.boxes == {"1": D("96708.00"), "2": D("10800.00"), "3": D("108708.00"), "4": D("6739.92"),
                       "5": D("108708.00"), "6": D("1576.32")}
    assert w.box12 == {"C": D("108.00"), "D": D("12000.00"), "DD": D("19200.00"), "W": D("4800.00")}
    assert w.box13_retirement_plan
    assert "7,800.00" in w.box14[0] and "Sec. 125" in w.box14[0]
    assert w.other_taxes == {"IL State Income Tax": D("4320.00")}


def test_w2_boxes_adp_caps_and_box10():
    w = estimate_w2(stub("paystub_adp.txt", "spouse"))
    assert w.periods_remaining == 20
    assert w.boxes["1"] == D("187975.54")                    # 210,833.42 − 13,113.08 − 9,744.80
    assert w.boxes["3"] == D("176100.00")                    # capped at the 2025 wage base
    assert w.boxes["4"] == LIMITS[2025].ss_max_tax == D("10918.20")
    assert w.boxes["5"] == D("201088.62")
    assert w.boxes["6"] == D("2925.58")                      # includes 0.9% over $200k
    assert w.boxes["10"] == D("4999.80")
    assert w.box12 == {"D": D("13113.08")}
    assert any("not derivable" in n for n in w.notes)        # no employer share → no DD


def test_limits_table():
    assert limits_for(2025).ss_wage_base == D("176100") and limits_for(2026).ss_wage_base == D("184500")
    assert limits_for(2025).elective_deferral == D("23500") and limits_for(2026).elective_deferral == D("24500")
    assert (limits_for(2026).hsa_self, limits_for(2026).hsa_family) == (D("4400"), D("8750"))
    assert (limits_for(2025).health_fsa, limits_for(2026).health_fsa) == (D("3300"), D("3400"))
    assert (limits_for(2025).dcap, limits_for(2026).dcap) == (D("5000"), D("7500"))
    assert limits_for(2031) is LIMITS[max(LIMITS)]


# --------------------------------------------------------------------------- household signals
def test_excess_social_security_across_employers():
    r = household()
    (s,) = by_id(r.signals, "excess_social_security")
    assert s.owner == "you" and s.amount == D("2517.72")      # 6,739.92 + 6,696.00 − 10,918.20
    assert "Schedule 3" in s.citation and s.severity == "money"
    assert not by_id(r.signals, "excess_social_security", "spouse")   # one employer: no credit signal


def test_402g_room_and_excess():
    r = household()
    (over,) = by_id(r.signals, "excess_deferral", "you")      # 12,000 pre-tax + 12,000 Roth
    assert over.amount == D("500.00") and "Notice 2024-80" in over.citation
    (room,) = by_id(r.signals, "deferral_room", "spouse")
    assert room.amount == D("10386.92")                       # 23,500 − 13,113.08


def test_hsa_fsa_conflict_and_limited_fsa():
    r = household()
    (c,) = by_id(r.signals, "hsa_fsa_conflict")
    assert "Pub. 969" in c.citation
    w = estimate_w2(stub("paystub_rippling.pdf"))
    w.health_fsa_type = "limited"
    assert not by_id(household_signals([w], 2025), "hsa_fsa_conflict")
    (h,) = by_id(household_signals([w], 2025), "hsa_over_limit")
    assert h.amount == D("500.00")                            # 4,800 vs 4,300 self-only


def test_dcap_warning_needs_the_flag():
    assert not by_id(household().signals, "dcap_no_earned_income")
    (s,) = by_id(household(spouse_no_earned_income=True).signals, "dcap_no_earned_income")
    assert s.amount == D("4999.80")


def test_additional_medicare_by_filing_status():
    r = household()
    assert r.filing_status == "mfj"
    (s,) = by_id(r.signals, "additional_medicare_not_withheld")
    # (108,708 + 108,000 + 201,088.62 − 250,000) × 0.9% − 0.9% × 1,088.62 withheld
    assert s.amount == D("1500.37") and "Form 8959" in s.citation and s.severity == "info"
    you_only = analyze_paystubs([(n, (FX / n).read_bytes(), "you")
                                 for n in ("paystub_rippling.pdf", "paystub_generic.txt")])
    (single,) = by_id(you_only.signals, "additional_medicare_not_withheld")
    assert you_only.filing_status == "single" and single.amount == D("150.37")   # (216,708 − 200,000) × 0.9%


def test_latest_stub_per_employer_wins():
    old = (FX / "paystub_generic.txt").read_text(encoding="utf-8").replace("2025-06", "2025-05").replace("9,000.00   54,000.00",
                                                                                          "9,000.00   45,000.00")
    r = analyze_paystubs([("old.txt", old.encode(), "you"),
                          ("paystub_generic.txt", (FX / "paystub_generic.txt").read_bytes(), "you")])
    assert len(r.w2) == 1 and r.w2[0].source == "paystub_generic.txt"
    assert any("older stub" in w for w in r.warnings)


def test_bad_inputs_rejected():
    with pytest.raises(ValueError):
        analyze_paystubs([("a.txt", b"x", "partner")])
    with pytest.raises(ValueError):
        analyze_paystubs([], frequency="daily")


# --------------------------------------------------------------------------- privacy: masking
def test_ssn_and_account_numbers_masked_everywhere(tmp_path):
    r = household()
    blob = json.dumps(to_jsonable(r))
    from ledgerstitch.reports.paystub_export import export_paystub_csv, export_paystub_xlsx
    csv_text = export_paystub_csv(r, tmp_path / "w2.csv").read_text(encoding="utf-8")
    from openpyxl import load_workbook
    wb = load_workbook(export_paystub_xlsx(r, tmp_path / "w2.xlsx"))
    cells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    for out in (blob, csv_text, cells):
        for secret in SECRETS:
            assert secret not in out
    text = (FX / "paystub_rippling.pdf").read_bytes()
    st = parse_paystub("Jane 900-45-6789.pdf", text)
    assert "900-45-6789" not in st.source


# --------------------------------------------------------------------------- export
def test_csv_export_is_labelled(tmp_path):
    from ledgerstitch.reports.paystub_export import export_paystub_csv
    rows = list(csv.reader(io.StringIO(export_paystub_csv(household(), tmp_path / "w2.csv").read_text(encoding="utf-8"))))
    assert rows[0] == ["ESTIMATE — not issued by employer"]
    assert rows[2] == ["section", "owner", "employer", "item", "amount", "note"]
    box1 = [r for r in rows[3:] if r[1] == "you" and r[2] == "Acme Corp" and r[3].startswith("Box 1 ")]
    assert box1[0][4] == "96708.00"
    assert any(r[0] == "Signal" and r[4] == "2517.72" for r in rows[3:])


def test_xlsx_export_is_labelled(tmp_path):
    from openpyxl import load_workbook

    from ledgerstitch.reports.paystub_export import export_paystub_xlsx
    wb = load_workbook(export_paystub_xlsx(household(), tmp_path / "w2.xlsx"))
    assert wb.sheetnames == ["Estimated W-2", "Signals", "Pay stubs"]
    for ws in wb:
        assert ws["A1"].value == ESTIMATE_LABEL


# --------------------------------------------------------------------------- CLI
def test_cli_paystub(capsys, tmp_path):
    out = tmp_path / "w2.csv"
    rc = cli_main(["paystub", str(FX / "paystub_rippling.pdf"), str(FX / "paystub_generic.txt"),
                   "--spouse", str(FX / "paystub_adp.txt"), "--out", str(out)])
    text = capsys.readouterr().out
    assert rc == 0 and ESTIMATE_LABEL in text and "excess Social Security" in text
    assert "96,708.00" in text and out.read_text(encoding="utf-8").startswith(ESTIMATE_LABEL)
    for secret in SECRETS:
        assert secret not in text


def test_cli_paystub_json_final(capsys):
    rc = cli_main(["paystub", str(FX / "paystub_rippling.pdf"), "--final", "--json"])
    data = json.loads(capsys.readouterr().out)
    assert rc == 0 and data["w2"][0]["final"] is True and data["w2"][0]["boxes"]["2"] == "4950.00"
    assert data["label"] == ESTIMATE_LABEL


def test_cli_paystub_missing_file(capsys, tmp_path):
    assert cli_main(["paystub", str(tmp_path / "nope.pdf")]) == 1
